from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path(r"D:\Codex_ws\Project\erounsolution_project")
REQ_DIR = ROOT / "docs" / "erounsolution_requests" / "requests"
DOCX_DIR = REQ_DIR / "docx"
XLSX_DIR = REQ_DIR / "xlsx"
IMG_DIR = REQ_DIR / "images"
SOURCE_IMAGE = IMG_DIR / "2026-06-16_ut_system_configuration_source.png"

DOCX_OUT = DOCX_DIR / "2026-06-16_UT시스템_구성_문의서.docx"
XLSX_OUT = XLSX_DIR / "이로운솔루션_문의사항_관리표.xlsx"


ITEMS = [
    ["UT 시스템 구성", "전체 UT 시스템 구성 범위", "스캐너, 소프트웨어, 무선 통신, 펌프, 물탱크, 분사기, 기타 구성품 포함 여부"],
    ["스캐너", "전원 사양", "정격 전압, 소비 전력, 전원 커넥터"],
    ["스캐너", "도면", "외형 치수, 장착부, 인터페이스 위치"],
    ["스캐너", "입출력 사양", "제어 입력, 상태 출력, 트리거, 안전 신호"],
    ["스캐너", "필요 스펙", "속도, 정밀도, 하중, 방수/방진, 사용 환경"],
    ["Software", "소프트웨어 구성 및 연동 방식", "운영 PC 필요 여부, SDK/API 제공 여부, 데이터 저장 포맷"],
    ["무선 통신 사양", "통신 방식 확인", "5G 또는 6G 필요 여부, Wi-Fi/이더넷 가능 여부, 지연시간 요구"],
    ["펌프", "모델명", "제조사, 모델, 공급 가능 여부"],
    ["펌프", "전원 사양", "정격 전압, 소비 전력, 전원 커넥터"],
    ["펌프", "입출력 사양", "제어 입력, 상태 출력, 유량 제어 가능 여부"],
    ["펌프", "필요 압력", "최소/최대 압력, 권장 운전 압력"],
    ["물탱크", "용량", "탱크 용량, 보충 방식, 설치 위치 제약"],
    ["분사기", "분사기 모델명", "제조사, 모델, 노즐 사양"],
    ["기타 필요 구성품 및 전원사양", "추가 구성품 목록", "밸브, 필터, 유량계, 압력센서, 케이블, 전원장치 등"],
]


def set_run_font(run, size=11, bold=False, color=None):
    run.font.name = "Malgun Gothic"
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_font(paragraph, size=11):
    for run in paragraph.runs:
        set_run_font(run, size=size)


def add_heading(doc, text, level=1):
    p = doc.add_heading("", level=level)
    run = p.add_run(text)
    set_run_font(run, size=16 if level == 1 else 13, bold=True, color="2E74B5" if level == 1 else "1F4D78")
    return p


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def style_table(table, header=True):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row_idx, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                set_paragraph_font(p, 9.5)
            if header and row_idx == 0:
                shade_cell(cell, "F2F4F7")
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.bold = True


def build_docx():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = doc.styles
    styles["Normal"].font.name = "Malgun Gothic"
    styles["Normal"].font.size = Pt(10.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("이로운 솔루션 측 문의 사항\nUT 시스템 구성")
    set_run_font(run, size=20, bold=True, color="0B2545")

    meta = doc.add_table(rows=4, cols=2)
    meta.columns[0].width = Inches(1.3)
    meta.columns[1].width = Inches(5.2)
    meta_data = [
        ["작성일", "2026-06-16"],
        ["프로젝트", "소형 원자로 비파괴 검사 자동화 프로젝트"],
        ["문의 대상", "이로운 솔루션"],
        ["상태", "문의 준비"],
    ]
    for row, values in zip(meta.rows, meta_data):
        for cell, value in zip(row.cells, values):
            cell.text = value
    style_table(meta, header=False)
    for row in meta.rows:
        shade_cell(row.cells[0], "E8EEF5")
        for run in row.cells[0].paragraphs[0].runs:
            run.bold = True

    add_heading(doc, "1. 문의 목적", 1)
    p = doc.add_paragraph(
        "SMR 비파괴 검사 자동화 시스템 설계를 위해 UT 시스템 구성품의 모델명, 전원 사양, "
        "입출력 사양, 통신 방식, 필요 스펙을 확인하고자 합니다."
    )
    set_paragraph_font(p)

    add_heading(doc, "2. 원본 문의 이미지", 1)
    if SOURCE_IMAGE.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(SOURCE_IMAGE), width=Inches(4.2))
        cap = doc.add_paragraph("그림 1. 사용자 제공 문의 항목 원본")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_font(cap, size=9)

    add_heading(doc, "3. 문의 항목", 1)
    table = doc.add_table(rows=1, cols=4)
    hdr = table.rows[0].cells
    headers = ["구분", "문의 내용", "확인 요청 세부사항", "회신"]
    for cell, value in zip(hdr, headers):
        cell.text = value
    for category, question, detail in ITEMS:
        cells = table.add_row().cells
        cells[0].text = category
        cells[1].text = question
        cells[2].text = detail
        cells[3].text = ""
    widths = [1.25, 1.55, 2.85, 0.85]
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)
    style_table(table)

    add_heading(doc, "4. 확인 요청 문안", 1)
    requests = [
        "UT 시스템 전체 구성품 목록을 공유 부탁드립니다.",
        "스캐너의 전원 사양, 도면, 입출력 사양, 필요 스펙을 공유 부탁드립니다.",
        "UT 관련 소프트웨어 구성, 운영 방식, SDK/API 제공 여부, 데이터 저장 포맷을 확인 부탁드립니다.",
        "무선 통신은 5G 또는 6G가 필수인지, Wi-Fi 또는 유선 이더넷 사용이 가능한지 확인 부탁드립니다.",
        "펌프의 모델명, 전원 사양, 입출력 사양, 필요 압력을 공유 부탁드립니다.",
        "물탱크 용량과 설치 제약 조건을 확인 부탁드립니다.",
        "분사기 모델명 및 노즐 사양을 공유 부탁드립니다.",
        "기타 필요한 구성품과 각 구성품의 전원 사양을 공유 부탁드립니다.",
    ]
    for item in requests:
        p = doc.add_paragraph(style=None)
        p.style = doc.styles["List Number"]
        p.add_run(item)
        set_paragraph_font(p)

    add_heading(doc, "5. 회신 후 반영 예정 문서", 1)
    for item in [
        "docs/02_requirements.md",
        "docs/03_system_architecture.md",
        "docs/04_inspection_workflow.md",
        "docs/05_software_plan.md",
        "docs/07_data_management.md",
        "docs/erounsolution_requests/README.md",
    ]:
        p = doc.add_paragraph(style=None)
        p.style = doc.styles["List Bullet"]
        p.add_run(item)
        set_paragraph_font(p)

    doc.save(DOCX_OUT)


def build_xlsx():
    wb = Workbook()
    ws = wb.active
    ws.title = "문의사항 관리"

    ws["A1"] = "이로운 솔루션 문의사항 관리표"
    ws["A1"].font = Font(name="맑은 고딕", size=16, bold=True, color="0B2545")
    ws.merge_cells("A1:I1")

    ws["A2"] = "프로젝트"
    ws["B2"] = "소형 원자로 비파괴 검사 자동화 프로젝트"
    ws["D2"] = "작성일"
    ws["E2"] = "2026-06-16"
    ws["G2"] = "상태"
    ws["H2"] = "문의 준비"

    headers = ["ID", "구분", "문의 내용", "확인 요청 세부사항", "우선순위", "상태", "담당", "회신 내용", "반영 문서"]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col, value=header)
        cell.font = Font(name="맑은 고딕", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4D78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for idx, (category, question, detail) in enumerate(ITEMS, start=1):
        row = 4 + idx
        values = [
            f"ERS-Q-{idx:03d}",
            category,
            question,
            detail,
            "높음" if category in ["UT 시스템 구성", "스캐너", "펌프"] else "중간",
            "문의 준비",
            "미정",
            "",
            "",
        ]
        for col, value in enumerate(values, start=1):
            ws.cell(row=row, column=col, value=value)

    widths = [12, 20, 24, 44, 12, 14, 14, 36, 32]
    for col_idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in ws.iter_rows(min_row=2, max_row=4 + len(ITEMS), min_col=1, max_col=9):
        for cell in row:
            cell.font = cell.font.copy(name="맑은 고딕", size=10)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = border
    for row in range(5, 5 + len(ITEMS)):
        ws.row_dimensions[row].height = 42
    ws.row_dimensions[4].height = 30

    dv_status = DataValidation(type="list", formula1='"문의 준비,문의 완료,회신 수신,검토중,반영 완료,보류"', allow_blank=False)
    ws.add_data_validation(dv_status)
    dv_status.add(f"F5:F{4 + len(ITEMS)}")

    dv_priority = DataValidation(type="list", formula1='"높음,중간,낮음"', allow_blank=False)
    ws.add_data_validation(dv_priority)
    dv_priority.add(f"E5:E{4 + len(ITEMS)}")

    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:I{4 + len(ITEMS)}"

    img_ws = wb.create_sheet("원본 이미지")
    img_ws["A1"] = "사용자 제공 문의 항목 원본"
    img_ws["A1"].font = Font(name="맑은 고딕", size=14, bold=True, color="0B2545")
    if SOURCE_IMAGE.exists():
        img = XLImage(str(SOURCE_IMAGE))
        img.width = 470
        img.height = 528
        img_ws.add_image(img, "A3")
    img_ws.column_dimensions["A"].width = 70
    img_ws.row_dimensions[3].height = 300

    wb.save(XLSX_OUT)


if __name__ == "__main__":
    DOCX_DIR.mkdir(parents=True, exist_ok=True)
    XLSX_DIR.mkdir(parents=True, exist_ok=True)
    build_docx()
    build_xlsx()
    print(DOCX_OUT)
    print(XLSX_OUT)
