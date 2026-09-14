"""검사 운영 기록을 '데이터 저장 위치' 아래에 파일로 남긴다.

항목마다 폴더를 따로 두고, 그 안을 연도/월 폴더로 나누고, 파일 이름은
날짜로 시작한다.

  <데이터 저장 위치>/
    작업기록/2026/09/20260914_작업기록.xlsx
        구간(격자) 하나당 한 줄. 사람이 보고 거르는 표라 엑셀로 둔다.
    스캔좌표/2026/09/20260914_153012_jb00000001_1A.txt
        구간 하나당 파일 하나. ㄹ자 스캔 중 원점 기준 좌표를 10 Hz 로.
    알람이벤트/2026/09/20260914_알람이벤트.txt
        하루 파일 하나. 진행 알림·알람·장애·해제·알림 메시지.
    통신기록/2026/09/20260914_통신기록.txt
        하루 파일 하나. ERUT·MC MQTT 송수신 원문.

계속 덧붙는 기록(좌표·알람·통신)은 **탭으로 나눈 텍스트**다. 한 줄씩
바로 파일 끝에 붙이므로 프로그램이 도중에 죽어도 그때까지는 남는다 —
알람 기록은 바로 그런 순간에 필요하다. 엑셀에서 열 때는 탭 구분으로
열면 표가 된다. 작업기록만 엑셀 파일이다(줄이 적고 보고용이라).

기록 실패가 검사를 멈추면 안 된다. 쓰다가 실패하면 `problem` 으로
한 번 알리고 넘어간다. 통신 기록은 MQTT 수신 스레드에서도 불리므로
파일 쓰기는 모두 잠금 안에서 한다.
"""

from __future__ import annotations

import os
import re
import tempfile
import threading
from collections.abc import Callable
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

from PyQt6.QtCore import QObject, QTimer, pyqtSignal

JOBS = "작업기록"
SCAN = "스캔좌표"
EVENTS = "알람이벤트"
COMMS = "통신기록"
CATEGORIES = (JOBS, SCAN, EVENTS, COMMS)

#: 로봇 태스크 상태(레지스터 290) 중 ㄹ자 스캔 중. 이때의 좌표만 남긴다.
ROBOT_STATE_SCANNING = 6

#: 통신 기록에서 뺄 토픽. 좌표 스트림은 스캔좌표 기록에 이미 있고,
#: 하트비트는 내용이 없다 — 넣으면 하루에 수십만 줄이 된다.
COMMS_SKIP_TOPICS = ("doosan/robot/tcp", "Heartbeat/robot")

JOB_COLUMNS = (
    ("날짜", 11), ("시작", 10), ("종료", 10), ("소요[s]", 8), ("출처", 7),
    ("job_id", 14), ("구간", 7), ("결과", 8), ("비고", 30),
    ("호 길이[mm]", 11), ("격자 높이[mm]", 12), ("격자간 겹침[mm]", 14),
    ("반지름[mm]", 10), ("두께[mm]", 9), ("프로브", 11), ("태스크", 9),
    ("좌표 수", 8), ("스캔 좌표 파일", 42),
)

#: 새 텍스트 파일 맨 앞에 붙이는 UTF-8 표시. 없으면 Windows 엑셀이 한글을
#: CP949 로 읽어 깨뜨린다(메모장은 없어도 읽는다). 이어 쓰기(a)에서는
#: 파일을 처음 만들 때 한 번만 쓴다.
_BOM = "\ufeff"

_DATE_PREFIX = re.compile(r"^(\d{8})_")
_WIN_DRIVE = re.compile(r"^([A-Za-z]):[\\/]*(.*)$")
_WSL_UNC = re.compile(r"^[\\/]{2}(?:wsl\.localhost|wsl\$)[\\/][^\\/]+(.*)$", re.I)


def _in_wsl() -> bool:
    if os.environ.get("WSL_DISTRO_NAME"):
        return True
    try:
        with open("/proc/sys/kernel/osrelease", encoding="utf-8") as f:
            return "microsoft" in f.read().lower()
    except OSError:
        return False


def resolve_data_root(value: str, in_wsl: bool | None = None) -> Path:
    """설정값(화면 표기)을 이 프로그램이 쓸 수 있는 폴더 경로로 바꾼다.

    * `SMR_DATA_DIR` 환경변수가 있으면 그것이 우선이다(시험·현장 고정용).
    * "D:/SMR/Data" 같은 Windows 경로는 WSL 에서 /mnt/d/SMR/Data 로 쓴다.
      WSL 이 아닌 리눅스에서는 쓸 수 없으므로 ~/SMR/Data 로 대신한다.
    * "//wsl.localhost/<배포판>/home/..." (Windows 창에서 WSL 폴더를 고른
      경우)는 /home/... 로 바꾼다.
    """
    override = os.environ.get("SMR_DATA_DIR", "").strip()
    if override:
        return Path(override).expanduser()
    value = (value or "").strip()
    fallback = Path.home() / "SMR" / "Data"
    if not value:
        return fallback
    unc = _WSL_UNC.match(value)
    if unc:
        rest = unc.group(1).replace("\\", "/")
        return Path(rest or "/")
    drive = _WIN_DRIVE.match(value)
    if drive:
        if in_wsl is None:
            in_wsl = _in_wsl()
        if not in_wsl:
            return fallback
        rest = drive.group(2).replace("\\", "/").strip("/")
        return Path("/mnt") / drive.group(1).lower() / rest
    return Path(value).expanduser()


def _stamp(now: datetime) -> str:
    return now.strftime("%Y-%m-%d %H:%M:%S.") + f"{now.microsecond // 1000:03d}"


def _clean(text: Any) -> str:
    """탭 구분 한 줄에 넣을 수 있게 탭·줄바꿈을 공백으로 바꾼다."""
    return re.sub(r"[\t\r\n]+", " ", str(text)).strip()


def _safe_name(text: str) -> str:
    return re.sub(r'[\\/:*?"<>|\s]+', "_", text).strip("_") or "-"


class DataRecorder(QObject):
    """네 가지 운영 기록을 파일로 남긴다."""

    #: 기록을 못 남겼을 때 사람이 볼 문장. 같은 문제는 한 번만 낸다.
    problem = pyqtSignal(str)

    #: 보존 기간 정리 주기 [ms]. 켜 둔 채 날이 바뀌어도 지워지게.
    PURGE_INTERVAL_MS = 6 * 60 * 60 * 1000
    #: 엑셀이 열려 있어 못 쓴 작업기록을 다시 써 보는 주기 [ms].
    RETRY_INTERVAL_MS = 60 * 1000

    def __init__(self, root_value: str = "", retention_days: int = 365,
                 clock: Callable[[], datetime] = datetime.now,
                 parent: QObject | None = None) -> None:
        super().__init__(parent)
        self._clock = clock
        self._lock = threading.RLock()
        self._root = resolve_data_root(root_value)
        self._retention_days = max(int(retention_days), 1)
        self._reported: set[str] = set()
        # 작업 맥락 — begin_job 이 채운다.
        self._job: dict[str, Any] = {}
        # 진행 중인 구간.
        self._cell: dict[str, Any] | None = None
        self._robot_state = 0
        self._scan_handle = None
        # 엑셀이 열려 있어 못 쓴 작업기록 줄. {파일: [줄, ...]}
        self._pending_rows: dict[Path, list[list[Any]]] = {}

        self._purge_timer = QTimer(self)
        self._purge_timer.setInterval(self.PURGE_INTERVAL_MS)
        self._purge_timer.timeout.connect(self.purge_old)
        self._purge_timer.start()
        self._retry_timer = QTimer(self)
        self._retry_timer.setInterval(self.RETRY_INTERVAL_MS)
        self._retry_timer.timeout.connect(self._flush_pending_rows)

    # ------------------------------------------------------------ 설정
    @property
    def root(self) -> Path:
        return self._root

    def set_root(self, value: str) -> Path:
        """저장 위치를 바꾼다. 진행 중인 좌표 파일은 원래 자리에서 끝낸다."""
        with self._lock:
            root = resolve_data_root(value)
            if root != self._root:
                self._root = root
                self._reported.discard("root")
            return self._root

    def set_retention_days(self, days: Any) -> None:
        try:
            self._retention_days = max(int(float(days)), 1)
        except (TypeError, ValueError):
            return
        self.purge_old()

    def path_for(self, category: str, filename: str, when: datetime | None = None) -> Path:
        """<저장 위치>/<항목>/<연도>/<월>/<파일>."""
        when = when or self._clock()
        return self._root / category / f"{when:%Y}" / f"{when:%m}" / filename

    # ------------------------------------------------------------ 작업기록
    def begin_job(self, source: str, job_id: str, plan: Any = None,
                  nosensor: bool | None = None) -> None:
        """작업(순회) 하나의 맥락을 잡는다. 구간 줄마다 같이 적힌다."""
        now = self._clock()
        if not job_id:
            job_id = f"{source}-{now:%Y%m%d-%H%M%S}"
        probes = int(getattr(plan, "eoat_probes", 0) or 0)
        self._job = {
            "source": source,
            "job_id": job_id,
            "width": getattr(plan, "cell_width", ""),
            "height": getattr(plan, "cell_height", ""),
            "overlap": getattr(plan, "pitch_y", ""),
            "radius": getattr(plan, "radius", ""),
            "thickness": getattr(plan, "thickness", ""),
            "probe": {5: "5축 십자", 8: "8축 직사각"}.get(probes, "-"),
            "task": "" if nosensor is None else ("논센서판" if nosensor else "센서판"),
        }
        self.log_event("정보", f"작업 기록 시작: {source} {job_id}")

    def cell_started(self, label: str) -> None:
        """구간 스캔을 시작했다. 이전 구간이 열려 있으면 중단으로 닫는다."""
        with self._lock:
            if self._cell is not None and self._cell["label"] != label:
                self._finish_cell("중단")
            if self._cell is not None:          # 같은 구간 재시작(일시정지 후)
                return
            self._cell = {"label": label, "start": self._clock(), "points": 0,
                          "scan_path": None, "notes": []}

    def cell_finished(self, label: str, result: str = "완료") -> None:
        with self._lock:
            if self._cell is None or self._cell["label"] != label:
                return
            self._finish_cell(result)

    def job_stopped(self, result: str = "중단") -> None:
        """정지·중단으로 순회가 끝났다. 열린 구간을 그 결과로 닫는다."""
        with self._lock:
            if self._cell is not None:
                self._finish_cell(result)

    def mark_finished(self, marked: list, failed: list) -> None:
        now = self._clock()
        job_id = self._job.get("job_id", "") or f"MARK-{now:%Y%m%d-%H%M%S}"
        row = [f"{now:%Y-%m-%d}", "", f"{now:%H:%M:%S}", "", "ERUT", job_id, "마킹",
               f"성공 {len(marked)} / 실패 {len(failed)}",
               _clean("실패: " + ", ".join(map(str, failed)) if failed else
                      "마킹: " + ", ".join(map(str, marked))),
               "", "", "", "", "", "", "", "", ""]
        self._append_job_row(row, now)

    def note(self, text: str) -> None:
        """진행 중인 구간의 비고에 붙인다(그 구간에서 난 알람 등)."""
        with self._lock:
            if self._cell is not None and text not in self._cell["notes"]:
                self._cell["notes"].append(_clean(text))

    def _finish_cell(self, result: str) -> None:
        cell, self._cell = self._cell, None
        self._close_scan_file()
        if cell is None:
            return
        end = self._clock()
        start: datetime = cell["start"]
        job = self._job
        scan_path: Path | None = cell["scan_path"]
        row = [
            f"{start:%Y-%m-%d}", f"{start:%H:%M:%S}", f"{end:%H:%M:%S}",
            round((end - start).total_seconds(), 1),
            job.get("source", ""), job.get("job_id", ""), cell["label"], result,
            "; ".join(cell["notes"]),
            job.get("width", ""), job.get("height", ""), job.get("overlap", ""),
            job.get("radius", ""), job.get("thickness", ""),
            job.get("probe", ""), job.get("task", ""),
            cell["points"],
            str(scan_path.relative_to(self._root)) if scan_path else "",
        ]
        self._append_job_row(row, start)

    def _append_job_row(self, row: list[Any], when: datetime) -> None:
        path = self.path_for(JOBS, f"{when:%Y%m%d}_{JOBS}.xlsx", when)
        with self._lock:
            self._pending_rows.setdefault(path, []).append(row)
        self._flush_pending_rows()

    def _flush_pending_rows(self) -> None:
        with self._lock:
            for path in list(self._pending_rows):
                rows = self._pending_rows[path]
                try:
                    _append_xlsx(path, rows)
                except PermissionError:
                    # Windows 에서 엑셀로 열어 두면 파일이 잠긴다. 버리지 않고
                    # 들고 있다가 닫히면 같이 쓴다.
                    self._report(f"xlsx:{path}",
                                 f"작업기록 파일이 다른 프로그램(엑셀)에서 열려 있어 "
                                 f"저장을 미뤘습니다 — 닫으면 자동으로 저장합니다: {path.name}")
                    self._retry_timer.start()
                    continue
                except OSError as exc:
                    self._report(f"xlsx:{path}", f"작업기록을 저장하지 못했습니다: {exc}")
                    self._retry_timer.start()
                    continue
                del self._pending_rows[path]
                self._reported.discard(f"xlsx:{path}")
            if not self._pending_rows:
                self._retry_timer.stop()

    # ------------------------------------------------------------ 스캔좌표
    def set_robot_state(self, state: int) -> None:
        self._robot_state = int(state)

    def scan_point(self, values: list) -> None:
        """원점 기준 좌표 한 개. ㄹ자 스캔 중(290 == 6)인 구간에서만 남긴다."""
        if len(values) < 6 or self._robot_state != ROBOT_STATE_SCANNING:
            return
        with self._lock:
            cell = self._cell
            if cell is None:
                return
            now = self._clock()
            try:
                if self._scan_handle is None:
                    self._open_scan_file(cell, now)
                elapsed = (now - cell["start"]).total_seconds()
                line = [_stamp(now), f"{elapsed:.3f}", cell["label"],
                        *(f"{float(v):.2f}" for v in values[:6])]
                self._scan_handle.write("\t".join(line) + "\n")
                self._scan_handle.flush()
                cell["points"] += 1
            except OSError as exc:
                self._close_scan_file()
                self._report("scan", f"스캔 좌표를 저장하지 못했습니다: {exc}")

    def _open_scan_file(self, cell: dict, now: datetime) -> None:
        job_id = _safe_name(str(self._job.get("job_id", "") or "job"))
        start: datetime = cell["start"]
        name = f"{start:%Y%m%d_%H%M%S}_{job_id}_{_safe_name(cell['label'])}.txt"
        path = self.path_for(SCAN, name, start)
        path.parent.mkdir(parents=True, exist_ok=True)
        new = not path.exists()
        handle = path.open("a", encoding="utf-8")
        if new:
            handle.write(_BOM)
        handle.write(
            "# 스캔 좌표 기록 — 원점(영점) 기준, Rz +180° 좌표계(X·Y 부호 반전), "
            "단위 mm · deg, ㄹ자 스캔 중(로봇 상태 6)인 값만\n"
            f"# job_id={self._job.get('job_id', '')}\t구간={cell['label']}\t"
            f"출처={self._job.get('source', '')}\t시작={start:%Y-%m-%d %H:%M:%S}\n"
            "시각\t경과[s]\t구간\tX[mm]\tY[mm]\tZ[mm]\tRx[deg]\tRy[deg]\tRz[deg]\n")
        self._scan_handle = handle
        cell["scan_path"] = path

    def _close_scan_file(self) -> None:
        handle, self._scan_handle = self._scan_handle, None
        if handle is not None:
            try:
                handle.close()
            except OSError:
                pass

    # ------------------------------------------------------------ 알람·이벤트
    def log_event(self, kind: str, message: str, code: str = "", level: str = "") -> None:
        """구분: 정보(진행 알림) / 알람 / 장애 / 해제 / 알림."""
        now = self._clock()
        self._append_line(
            EVENTS, now, "시각\t구분\t코드\t수준\t내용",
            [_stamp(now), kind, code, level, _clean(message)],
            title="# 알람·이벤트 기록 — 구분: 정보(진행 알림) / 알람 / 장애 / 해제 / 알림")
        if kind in ("알람", "장애"):
            self.note(f"{code} {message}".strip())

    # ------------------------------------------------------------ 통신
    def comms(self, direction: str, channel: str, topic: str, payload: Any) -> None:
        """MQTT 송수신 한 건. 어느 스레드에서 불러도 된다."""
        if topic in COMMS_SKIP_TOPICS:
            return
        now = self._clock()
        if isinstance(payload, (bytes, bytearray)):
            payload = payload.decode("utf-8", "replace")
        self._append_line(
            COMMS, now, "시각\t방향\t채널\t토픽\t내용",
            [_stamp(now), direction, channel, topic, _clean(payload)],
            title=f"# 통신 기록 — MQTT 송수신 원문 (제외: {', '.join(COMMS_SKIP_TOPICS)} — "
                  "좌표는 스캔좌표 기록에 있다)")

    # ------------------------------------------------------------ 공통
    def _append_line(self, category: str, now: datetime, header: str,
                     fields: list[str], title: str) -> None:
        path = self.path_for(category, f"{now:%Y%m%d}_{category}.txt", now)
        with self._lock:
            try:
                path.parent.mkdir(parents=True, exist_ok=True)
                new = not path.exists()
                with path.open("a", encoding="utf-8") as handle:
                    if new:
                        handle.write(_BOM + title + "\n" + header + "\n")
                    handle.write("\t".join(fields) + "\n")
                self._reported.discard(category)
            except OSError as exc:
                self._report(category, f"{category}을(를) 저장하지 못했습니다 ({self._root}): {exc}")

    def _report(self, key: str, message: str) -> None:
        if key in self._reported:
            return
        self._reported.add(key)
        self.problem.emit(message)

    def purge_old(self, today: date | None = None) -> list[Path]:
        """보존 기간이 지난 기록 파일을 지운다. 우리가 만든 이름만 건드린다.

        파일 이름 앞 8자리 날짜로 판단한다(수정 시각은 복사하면 바뀐다).
        비게 된 월·연도 폴더도 지운다.
        """
        today = today or self._clock().date()
        cutoff = today - timedelta(days=self._retention_days)
        removed: list[Path] = []
        with self._lock:
            for category in CATEGORIES:
                base = self._root / category
                if not base.is_dir():
                    continue
                for path in base.glob("*/*/*"):
                    m = _DATE_PREFIX.match(path.name)
                    if not m or path.suffix not in (".txt", ".xlsx") or not path.is_file():
                        continue
                    try:
                        day = datetime.strptime(m.group(1), "%Y%m%d").date()
                    except ValueError:
                        continue
                    if day < cutoff:
                        try:
                            path.unlink()
                            removed.append(path)
                        except OSError:
                            pass
                for month in sorted(base.glob("*/*"), reverse=True):
                    _rmdir_if_empty(month)
                for year in base.glob("*"):
                    _rmdir_if_empty(year)
        return removed

    def close(self) -> None:
        with self._lock:
            self._close_scan_file()
            self._flush_pending_rows()


def _rmdir_if_empty(path: Path) -> None:
    try:
        if path.is_dir() and not any(path.iterdir()):
            path.rmdir()
    except OSError:
        pass


def _append_xlsx(path: Path, rows: list[list[Any]]) -> None:
    """작업기록 엑셀에 줄을 덧붙인다. 임시 파일에 쓴 뒤 바꿔치기한다.

    쓰는 도중 꺼져도 원래 파일은 멀쩡하다. 파일이 엑셀로 열려 있으면
    (Windows 잠금) PermissionError 가 난다 — 부르는 쪽이 나중에 다시 쓴다.
    """
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        book = load_workbook(path)
        sheet = book.active
    else:
        book = Workbook()
        sheet = book.active
        sheet.title = JOBS
        sheet.append([name for name, _ in JOB_COLUMNS])
        head = PatternFill("solid", fgColor="DCE6F1")
        for cell, (_, width) in zip(sheet[1], JOB_COLUMNS):
            cell.font = Font(bold=True)
            cell.fill = head
            cell.alignment = Alignment(horizontal="center")
            sheet.column_dimensions[cell.column_letter].width = width
        sheet.freeze_panes = "A2"
    for row in rows:
        sheet.append(row)
    handle = tempfile.NamedTemporaryFile(dir=path.parent, prefix=path.stem,
                                         suffix=".tmp", delete=False)
    handle.close()
    try:
        book.save(handle.name)
        os.replace(handle.name, path)
    except BaseException:
        Path(handle.name).unlink(missing_ok=True)
        raise
