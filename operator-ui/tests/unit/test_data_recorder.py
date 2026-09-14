"""운영 기록(작업·스캔 좌표·알람이벤트·통신) 파일 저장 시험.

폴더 구조: <저장 위치>/<항목>/<연도>/<월>/<날짜로 시작하는 파일>.
"""

from datetime import date, datetime, timedelta
from pathlib import Path

import pytest
from openpyxl import load_workbook

from smr_operator_ui.services import GridPlan
from smr_operator_ui.services import data_recorder as dr
from smr_operator_ui.services.data_recorder import DataRecorder, resolve_data_root


class Clock:
    """시험용 시계. 호출할 때마다 조금씩 흐른다."""

    def __init__(self, start: datetime) -> None:
        self.now = start

    def __call__(self) -> datetime:
        self.now += timedelta(milliseconds=100)
        return self.now


@pytest.fixture
def rec(qtbot, tmp_path, monkeypatch):
    monkeypatch.delenv("SMR_DATA_DIR", raising=False)
    r = DataRecorder(root_value=str(tmp_path), retention_days=365,
                     clock=Clock(datetime(2026, 9, 14, 15, 30, 0)))
    yield r
    r.close()


def _plan() -> GridPlan:
    return GridPlan(column_count=3, row_count=4, cell_width=721, cell_height=500,
                    scan_overlap=0.0, pitch_x=20, pitch_y=20, radius=834.6,
                    thickness=10, eoat_probes=8)


def _lines(path: Path) -> list[str]:
    raw = path.read_text(encoding="utf-8")
    assert raw.startswith("\ufeff"), "엑셀이 UTF-8 로 읽게 BOM 으로 시작한다"
    assert raw.count("\ufeff") == 1, "이어 쓸 때 BOM 을 또 넣으면 안 된다"
    return raw[1:].splitlines()


# ---- 폴더·파일 이름 -------------------------------------------------------------
def test_each_record_has_its_own_folder_by_year_and_month(rec, tmp_path) -> None:
    rec.begin_job("ERUT", "jb00000001", _plan(), nosensor=False)
    rec.cell_started("1A")
    rec.set_robot_state(dr.ROBOT_STATE_SCANNING)
    rec.scan_point([0.0, 12.5, 0.0, 0.0, 0.0, 0.0])
    rec.cell_finished("1A")
    rec.log_event("알람", "센터 프로브 벽 접촉 실패")
    rec.comms("수신", "ERUT", "doosan/robot/req/start", b'{"content":{}}')

    files = sorted(p.relative_to(tmp_path).as_posix() for p in tmp_path.rglob("*") if p.is_file())
    assert files == [
        "스캔좌표/2026/09/20260914_153000_jb00000001_1A.txt",
        "알람이벤트/2026/09/20260914_알람이벤트.txt",
        "작업기록/2026/09/20260914_작업기록.xlsx",
        "통신기록/2026/09/20260914_통신기록.txt",
    ]


# ---- 스캔좌표 -------------------------------------------------------------------
def test_scan_points_are_kept_only_while_scanning_inside_a_cell(rec, tmp_path) -> None:
    rec.begin_job("MC", "jb7", _plan())
    rec.set_robot_state(dr.ROBOT_STATE_SCANNING)
    rec.scan_point([1, 2, 3, 4, 5, 6])          # 구간 밖 — 버린다
    rec.cell_started("2B")
    rec.set_robot_state(7)                       # 원점 대기 — 버린다
    rec.scan_point([9, 9, 9, 9, 9, 9])
    assert not (tmp_path / dr.SCAN).exists(), "좌표가 없으면 파일도 안 만든다"

    rec.set_robot_state(dr.ROBOT_STATE_SCANNING)
    rec.scan_point([0.0, 75.8, 0.0, 0.0, 0.0, 180.0])
    rec.scan_point([1.25, 80.0, 29.4, 0.1, -0.2, 179.9])
    rec.cell_finished("2B")

    (path,) = (tmp_path / dr.SCAN).rglob("*.txt")
    lines = _lines(path)
    assert lines[0].startswith("# 스캔 좌표 기록")
    assert "job_id=jb7" in lines[1] and "구간=2B" in lines[1]
    assert lines[2].split("\t") == ["시각", "경과[s]", "구간", "X[mm]", "Y[mm]",
                                    "Z[mm]", "Rx[deg]", "Ry[deg]", "Rz[deg]"]
    rows = [line.split("\t") for line in lines[3:]]
    assert len(rows) == 2
    assert rows[1][2:] == ["2B", "1.25", "80.00", "29.40", "0.10", "-0.20", "179.90"]


# ---- 작업기록 -------------------------------------------------------------------
def test_job_rows_go_to_one_daily_workbook(rec, tmp_path) -> None:
    rec.begin_job("ERUT", "jb00000001", _plan(), nosensor=True)
    rec.cell_started("1A")
    rec.set_robot_state(dr.ROBOT_STATE_SCANNING)
    rec.scan_point([0, 1, 2, 0, 0, 0])
    rec.log_event("장애", "좌측 프로브 벽 접촉 실패", code="E-PROBE-L", level="stop")
    rec.cell_finished("1A")
    rec.cell_started("1B")
    rec.job_stopped()                            # 작업자 정지·ERUT abort

    (book,) = (tmp_path / dr.JOBS).rglob("*.xlsx")
    sheet = load_workbook(book).active
    header = [c.value for c in sheet[1]]
    assert header == [name for name, _ in dr.JOB_COLUMNS]
    first = dict(zip(header, [c.value for c in sheet[2]]))
    assert first["출처"] == "ERUT" and first["job_id"] == "jb00000001"
    assert first["구간"] == "1A" and first["결과"] == "완료"
    assert "E-PROBE-L" in first["비고"], "구간 중 난 장애는 비고에 남는다"
    assert first["프로브"] == "8축 직사각" and first["태스크"] == "논센서판"
    assert first["호 길이[mm]"] == 721 and first["격자간 겹침[mm]"] == 20
    assert first["좌표 수"] == 1
    assert first["스캔 좌표 파일"].startswith("스캔좌표/2026/09/20260914_")
    second = dict(zip(header, [c.value for c in sheet[3]]))
    assert second["구간"] == "1B" and second["결과"] == "중단"


def test_marking_adds_a_row(rec, tmp_path) -> None:
    rec.mark_finished(["p1", "p2"], ["p3"])
    (book,) = (tmp_path / dr.JOBS).rglob("*.xlsx")
    row = [c.value for c in load_workbook(book).active[2]]
    assert row[6] == "마킹" and row[7] == "성공 2 / 실패 1"


def test_locked_workbook_is_written_later(rec, tmp_path, monkeypatch) -> None:
    """엑셀로 열려 있어 잠긴 파일은 버리지 않고 나중에 같이 쓴다."""
    real = dr._append_xlsx
    calls = {"n": 0}

    def locked_once(path, rows):
        calls["n"] += 1
        if calls["n"] == 1:
            raise PermissionError("locked")
        real(path, rows)

    monkeypatch.setattr(dr, "_append_xlsx", locked_once)
    problems: list[str] = []
    rec.problem.connect(problems.append)
    rec.begin_job("RCS", "", _plan())
    rec.cell_started("1A"); rec.cell_finished("1A")
    assert problems and "열려 있어" in problems[0]
    assert not list((tmp_path / dr.JOBS).rglob("*.xlsx"))

    rec._flush_pending_rows()                    # 재시도 타이머가 하는 일
    (book,) = (tmp_path / dr.JOBS).rglob("*.xlsx")
    assert load_workbook(book).active.max_row == 2


# ---- 알람이벤트·통신 --------------------------------------------------------------
def test_events_and_comms_are_tab_separated_daily_text(rec, tmp_path) -> None:
    rec.log_event("알림", "현재 로봇이 동작 중이므로\t홈 이동이 불가합니다.", code="M1001")
    rec.comms("발신", "ERUT", "erut/robot1/res", '{"code":409,\n"message":"BUSY"}')
    rec.comms("발신", "MC", "doosan/robot/tcp", '{"x":1}')          # 좌표 스트림 — 제외
    rec.comms("수신", "MC", "Heartbeat/robot", '{}')                 # 하트비트 — 제외

    (events,) = (tmp_path / dr.EVENTS).rglob("*.txt")
    assert _lines(events)[1] == "시각\t구분\t코드\t수준\t내용"
    fields = _lines(events)[2].split("\t")
    assert fields[1:] == ["알림", "M1001", "", "현재 로봇이 동작 중이므로 홈 이동이 불가합니다."]

    (comms,) = (tmp_path / dr.COMMS).rglob("*.txt")
    rows = [line.split("\t") for line in _lines(comms)[2:]]
    assert len(rows) == 1
    assert rows[0][1:4] == ["발신", "ERUT", "erut/robot1/res"]
    assert rows[0][4] == '{"code":409, "message":"BUSY"}'


# ---- 저장 위치 해석 ----------------------------------------------------------------
def test_windows_paths_map_to_wsl_mounts(monkeypatch) -> None:
    monkeypatch.delenv("SMR_DATA_DIR", raising=False)
    assert resolve_data_root("D:/SMR/Data", in_wsl=True) == Path("/mnt/d/SMR/Data")
    assert resolve_data_root("E:\\UT\\기록", in_wsl=True) == Path("/mnt/e/UT/기록")
    assert resolve_data_root("//wsl.localhost/Ubuntu-22.04/home/user/data",
                             in_wsl=True) == Path("/home/user/data")
    # 실제 우분투에서는 Windows 경로를 쓸 수 없다 — 홈 아래로 대신한다.
    assert resolve_data_root("D:/SMR/Data", in_wsl=False) == Path.home() / "SMR" / "Data"
    assert resolve_data_root("/srv/smr", in_wsl=False) == Path("/srv/smr")
    monkeypatch.setenv("SMR_DATA_DIR", "/tmp/fixed")
    assert resolve_data_root("D:/SMR/Data", in_wsl=True) == Path("/tmp/fixed")


# ---- 보존 기간 -------------------------------------------------------------------
def test_old_records_are_purged_and_empty_folders_removed(rec, tmp_path) -> None:
    old = tmp_path / dr.EVENTS / "2025" / "08" / "20250801_알람이벤트.txt"
    keep = tmp_path / dr.EVENTS / "2026" / "09" / "20260901_알람이벤트.txt"
    foreign = tmp_path / dr.EVENTS / "2025" / "07" / "메모.txt"   # 우리가 만든 파일 아님
    for path in (old, keep, foreign):
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text("x", encoding="utf-8")

    rec._retention_days = 30
    removed = rec.purge_old(today=date(2026, 9, 14))

    assert removed == [old]
    assert not old.exists()
    assert not (tmp_path / dr.EVENTS / "2025" / "08").exists(), "빈 월 폴더는 지운다"
    assert keep.exists()
    assert foreign.exists(), "이름이 날짜로 시작하지 않는 파일은 건드리지 않는다"
