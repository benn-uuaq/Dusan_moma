"""오류 로그·로그 파일·운전 모드 저장 화면 시험.

역할 나눔: 저장 위치·보존 기간은 시스템 설정에서만 바꾸고, 로그 파일은
기록 파일 관리, 오류 로그는 활성 오류 해제 + 알람이벤트 기록, 운전 모드는
작업 조건 묶음 저장·불러오기.
"""

from pathlib import Path

import pytest
from openpyxl import load_workbook

from smr_operator_ui.app import OperatorWindow
from smr_operator_ui.services import data_recorder as dr


@pytest.fixture
def window(qtbot, tmp_path, monkeypatch):
    monkeypatch.setenv("SMR_DATA_DIR", str(tmp_path / "data"))
    w = OperatorWindow(start_mqtt=False, start_ros=False)
    qtbot.addWidget(w)
    # 로봇 명령은 가로챈다(로봇 없이 알람이 쌓이지 않게).
    w.ros_status.call_command = lambda name: True
    w.ros_status.set_task_paths = lambda scan, mark: True
    yield w
    w.close()


def _texts(table, col):
    return [table.item(r, col).text() for r in range(table.rowCount())]


# ---- 오류 로그 -------------------------------------------------------------------
def test_error_log_shows_active_faults_and_records(window) -> None:
    window.erut_session.raise_error({"code": "E-PROBE-C", "message": "센터 프로브 벽 접촉 실패",
                                     "level": "stop", "recovery": "manual"})
    window.main_screen.show_activity("진행 알림 한 줄")
    window.navigate("errors")
    screen = window.screens["errors"]

    assert [screen.active_list.item(i).text() for i in range(screen.active_list.count())] == ["E-PROBE-C"]
    kinds = _texts(screen.table, 1)
    assert "장애" in kinds and "정보" not in kinds, "기본 보기는 진행 알림을 뺀다"
    screen.kind_combo.setCurrentIndex(1)            # 전체(진행 알림 포함)
    screen._load_table()
    assert "정보" in _texts(screen.table, 1)


def test_clearing_one_fault_logs_the_release(window) -> None:
    for code in ("E-PROBE-C", "E-PROBE-L"):
        window.erut_session.raise_error({"code": code, "message": code, "level": "stop",
                                         "recovery": "manual"})
    window.navigate("errors")
    screen = window.screens["errors"]
    screen.active_list.item(0).setSelected(True)
    screen._clear_selected()

    assert window.erut_session.error_codes() == ["E-PROBE-L"]
    events = window.data_recorder.read_events(screen.selected_day())
    assert ["해제", "E-PROBE-C"] in [[r[1], r[2]] for r in events]

    screen.clear_all_button.click()                 # 전체 해제 = 알람 리셋
    assert window.erut_session.error_codes() == []


def test_new_events_appear_live_and_export_to_excel(window, tmp_path) -> None:
    window.navigate("errors")
    screen = window.screens["errors"]
    before = screen.table.rowCount()
    window.cobot_manual_screen.add_alarm("시험 알람")
    assert screen.table.rowCount() == before + 1
    assert screen.table.item(0, 4).text() == "시험 알람", "최근 것이 맨 위"

    out = tmp_path / "export"
    screen.picker.chooser = lambda start: str(out)
    screen.export_button.click()
    (book,) = out.glob("*_오류로그.xlsx")
    rows = list(load_workbook(book).active.iter_rows(values_only=True))
    assert list(rows[0]) == screen.HEADERS and rows[1][4] == "시험 알람"


# ---- 로그 파일 -------------------------------------------------------------------
def test_log_files_list_export_and_delete(window, tmp_path) -> None:
    window.cobot_manual_screen.add_alarm("기록을 하나 만든다")
    window.data_recorder.comms("수신", "ERUT", "doosan/robot/req/query", "{}")
    window.navigate("logs")
    screen = window.screens["logs"]
    assert str(window.data_recorder.root) in screen.location_label.text()
    assert "보존 기간" in screen.location_label.text()
    assert sorted(_texts(screen.table, 1)) == sorted([dr.EVENTS, dr.COMMS])

    screen.table.selectRow(0)
    out = tmp_path / "copy"
    screen.picker.chooser = lambda start: str(out)
    screen.export_button.click()
    assert len(list(out.rglob("*.txt"))) == 1, "선택한 파일만 복사한다"

    screen.confirm = lambda text: False             # 취소하면 안 지운다
    screen.delete_button.click()
    assert screen.table.rowCount() == 2
    screen.confirm = lambda text: True
    screen.delete_button.click()
    assert screen.table.rowCount() == 1
    assert len(list(out.rglob("*.txt"))) == 1, "내보낸 사본은 그대로"


def test_log_files_filter_and_links(window) -> None:
    window.cobot_manual_screen.add_alarm("알람")
    window.data_recorder.comms("발신", "MC", "doosan/robot/job_state", "{}")
    window.navigate("logs")
    screen = window.screens["logs"]
    screen.kind_combo.setCurrentIndex(screen.kind_combo.findText(dr.COMMS))
    screen.refresh()
    assert _texts(screen.table, 1) == [dr.COMMS]

    opened = []
    screen.opener = opened.append
    screen.open_button.click()
    assert opened == [str(window.data_recorder.root)]

    screen.settings_button.click()                  # 위치·보존 기간은 시스템 설정에서 바꾼다
    assert window.stack.currentWidget() is window.screens["system"]


# ---- 운전 모드 저장 -------------------------------------------------------------------
def _modes(window, name="시험 모드"):
    screen = window.screens["modes"]
    screen.ask_name = lambda default: name
    screen.confirm = lambda text: True
    return screen


def test_mode_slot_saves_and_restores_job_conditions(window) -> None:
    saved = []
    real_save = window.settings_service.save
    window.settings_service.save = lambda scope, values: (saved.append((scope, dict(values))),
                                                           real_save(scope, values))
    screen = _modes(window)
    cobot = window.screens["cobot"]
    cobot.field("작업 속도").setValue(90)
    window._task_version = "dusan_v5"
    cobot.set_task_version("dusan_v5")
    window.main_screen.set_work_area(700.0, 450.0, 30.0, 15.0)

    screen._save_clicked()                           # 슬롯 1
    assert "시험 모드" in screen.slot_buttons[0].text()
    slots = [v for s, v in saved if s == "mode_slots"][-1]
    snap = slots["1"]
    assert snap["cobot"]["작업 속도"] == 90
    assert snap["robot_task"]["task_version"] == "dusan_v5"
    assert snap["work_area"]["width_mm"] == 700.0 and snap["work_area"]["overlap_mm"] == 15.0
    assert "UT 주소" not in snap["ut"] and "통신 포트" not in snap["ut"], "장비 연결값은 넣지 않는다"
    assert "태스크 판" not in snap["cobot"]

    # 다른 값으로 바꿔 놓고 불러오면 슬롯 값으로 돌아온다.
    cobot.field("작업 속도").setValue(80)
    cobot.set_task_version("dusan_v4"); window._task_version = "dusan_v4"
    window.main_screen.set_work_area(600.0, 400.0, 30.0, 0.0)
    screen._load_clicked()
    assert cobot.field("작업 속도").value() == 90
    assert cobot.task_version() == "dusan_v5" and window._task_version == "dusan_v5"
    assert window.main_screen.rect_view.work_area()[0] == 700.0
    assert "불러왔습니다" in screen.status_label.text()


def test_mode_slot_is_not_loaded_during_a_job(window) -> None:
    screen = _modes(window)
    screen._save_clicked()
    window._start_inspection()
    window.screens["cobot"].field("작업 속도").setValue(77)
    screen._load_clicked()
    assert window.screens["cobot"].field("작업 속도").value() == 77
    assert "작업 중" in screen.status_label.text()
    window._stop_inspection()


def test_mode_slot_clear_and_restore_on_start(window) -> None:
    screen = _modes(window)
    # 켰을 때 저장소에서 온 값으로 슬롯이 채워진다.
    window._apply_stored_settings("mode_slots", {"3": {"name": "야간", "saved_at": "2026-09-15 20:00",
                                                       "cobot": {"작업 속도": 100}}, "1": {}})
    assert "야간" in screen.slot_buttons[2].text()
    assert "비어 있음" in screen.slot_buttons[0].text()

    screen._save_clicked()
    assert screen.load_button.isEnabled()
    screen.clear_button.click()
    assert "비어 있음" in screen.slot_buttons[0].text()
    assert not screen.load_button.isEnabled()
    assert "야간" in screen.slot_buttons[2].text(), "다른 슬롯은 그대로"


def test_late_startup_load_does_not_wipe_a_fresh_slot(window) -> None:
    """켤 때 불러오기가 늦게 와도 방금 저장한 슬롯을 덮지 않는다(실제로 겪음)."""
    screen = _modes(window, "방금 저장")
    screen._save_clicked()
    window._apply_stored_settings("mode_slots", {})     # 늦게 도착한 옛 값
    assert "방금 저장" in screen.slot_buttons[0].text()
    assert screen.load_button.isEnabled()


# ---- 수동 제어 (차량) -------------------------------------------------------------------
def test_manual_jog_repeats_while_held_and_stops_on_release(window, qtbot) -> None:
    screen = window.screens["manual"]
    sent = []
    screen.manual_requested.connect(sent.append)
    fwd = screen.jog_buttons["cmd_mv_fwd"]
    assert not fwd.isEnabled(), "차량 제어가 더미면 잠겨 있다"

    screen.set_available(True)
    fwd.pressed.emit()
    qtbot.wait(screen.REPEAT_MS * 2 + 80)
    fwd.released.emit()
    assert sent[0] == {"cmd_mv_fwd": True}
    assert sent.count({"cmd_mv_fwd": True}) >= 2, "누르는 동안 다시 보낸다"
    assert sent[-1] == {}, "떼면 '명령 없음'으로 멈춘다"

    sent.clear()
    screen.lift_target.setValue(1200)
    screen.lift_button.click()
    assert sent == [{"cmd_mv_lift": 1, "lift_height": 1.2}]


def test_manual_interlock_while_the_robot_works(window) -> None:
    screen = window.screens["manual"]
    screen.set_available(True)
    window._on_robot_task_state(1)                  # 로봇 태스크 실행 중
    assert not screen.jog_buttons["cmd_mv_fwd"].isEnabled()
    assert not screen.outrigger_release_button.isEnabled()
    assert not screen.lift_button.isEnabled()
    assert screen.outrigger_set_button.isEnabled() and screen.stop_button.isEnabled()
    assert "인터락" in screen.notice.text()
    window._on_robot_task_state(3)
    assert screen.jog_buttons["cmd_mv_fwd"].isEnabled()


def test_manual_status_panel(window) -> None:
    screen = window.screens["manual"]
    screen.set_status({"state": "HOLD", "hold": "SET", "error_code": 0, "job_id": "RCS-0003",
                       "mv_dist": 0.7, "set_dist": 0.7, "speed": 0.0, "sen1_dist": 0.51,
                       "sen2_dist": 0.49, "x": 2.1, "yaw": 0.0, "lift_h": 0.96})
    assert screen.lift_value.text() == "960 mm"
    assert screen.status_blocks["hold"].value_label.text() == "고정"
    assert screen.status_blocks["error"].value_label.text() == "정상"
